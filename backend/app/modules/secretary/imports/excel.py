from openpyxl import load_workbook


# ==========================================================
# LEER EXCEL DE ESTUDIANTES
# ==========================================================

def leer_excel_estudiantes(archivo):


    libro = load_workbook(
        archivo
    )


    hoja = libro.active


    estudiantes = []


    # Saltamos encabezado
    for fila in hoja.iter_rows(
        min_row=2,
        values_only=True
    ):


        if not fila[0]:
            continue


        estudiante = {


            # Convertir DNI a texto
            "dni": str(int(fila[0])),


            "nombres": str(fila[1]).strip(),


            "apellidos": str(fila[2]).strip(),


            # Código siempre texto
            "student_code": str(fila[3]).strip(),


            # Teléfono texto
            "telefono": str(int(fila[4])),


            "direccion": str(fila[5]).strip()

        }


        estudiantes.append(
            estudiante
        )


    return estudiantes

# ==========================================================
# LEER EXCEL DE DOCENTES
# ==========================================================

def leer_excel_docentes(archivo):


    libro = load_workbook(
        archivo
    )


    hoja = libro.active


    docentes = []


    # Saltamos encabezado
    for fila in hoja.iter_rows(
        min_row=2,
        values_only=True
    ):


        if not fila[0]:
            continue


        docente = {


            # DNI como texto
            "dni": str(int(fila[0])),


            "nombres": str(fila[1]).strip(),


            "apellidos": str(fila[2]).strip(),


            # Código docente
            "teacher_code": str(fila[3]).strip(),


            # Teléfono texto
            "telefono": str(int(fila[4])),


            "direccion": str(fila[5]).strip()

        }


        docentes.append(
            docente
        )


    return docentes